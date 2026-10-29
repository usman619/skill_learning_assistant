import os
from typing import Any, Dict, List, Optional
from app.agent_coordinator import AgentCoordinator
from app.logging.audit_log import AuditLogger
from app.memory.database_operations import DatabaseOperations
from app.memory.migrations import init_db
from app.policies.approval_rules import check_approval_required, detect_untrusted_instruction
from app.tools.evidence_evaluator import evaluate_practice_evidence
from app.tools.memory_retrieval import retrieve_learner_state
from app.tools.plan_reviser import revise_learning_plan
from app.tools.practice_recorder import record_practice_outcome
from app.tools.resource_search import search_learning_resources


class WorkflowController:
    def __init__(self, db_path: str):
        self.db_path = db_path
        init_db(db_path)
        self.db = DatabaseOperations(db_path)
        self.logger = AuditLogger(db_path)
        self.coordinator = AgentCoordinator(self.logger)

    def onboard_learner(
        self,
        learner_id: str,
        goal: str,
        current_level: str,
        hours_per_week: float,
        deadline: str,
        simulate_search_failure: bool = False,
    ) -> Dict[str, Any]:
        # G1: Save checkpoint before starting onboarding workflow
        chk_id = self.db.save_checkpoint(
            learner_id,
            "onboard_start",
            {"goal": goal, "level": current_level, "hours": hours_per_week, "deadline": deadline},
        )

        profile = self.db.create_or_update_profile(
            learner_id=learner_id,
            goal=goal,
            current_level=current_level,
            hours_per_week=hours_per_week,
            deadline=deadline,
        )
        self.logger.log_event(learner_id, "ONBOARDING", "create_or_update_profile", "committed", profile)

        # G3: Record onboarding as a long-term episode summary
        self.db.add_episode_summary(
            learner_id=learner_id,
            episode_type="preference",
            summary=f"Onboarded with goal '{goal}', level '{current_level}', {hours_per_week} hrs/week, deadline {deadline}.",
            metadata={"profile_version": profile["version"], "hours_per_week": hours_per_week},
        )

        resources = search_learning_resources(
            self.db,
            self.logger,
            learner_id=learner_id,
            skill=goal,
            topic="load csv inspect columns mean",
            level=current_level,
            simulate_failure=simulate_search_failure,
        )

        state = retrieve_learner_state(self.db, self.logger, learner_id)
        proposal = self.coordinator.recommend_next_action(
            learner_id=learner_id,
            state=state,
            latest_assessment=None,
            resources=resources,
        )

        plan = revise_learning_plan(
            self.db,
            self.logger,
            learner_id=learner_id,
            proposal=proposal,
            due_date=deadline,
        )

        # G3: Record initial commitment in long-term episode memory
        self.db.add_episode_summary(
            learner_id=learner_id,
            episode_type="commitment",
            summary=f"Plan v{plan['version']} assigned: '{proposal.recommended_task}' ({proposal.estimated_minutes} min).",
            metadata={"plan_id": plan["plan_id"], "plan_version": plan["version"]},
        )

        # G1: Mark onboarding checkpoint completed
        self.db.complete_checkpoint(chk_id)

        return {
            "profile": profile,
            "resources": resources,
            "proposal": proposal.model_dump(),
            "plan": plan,
        }

    def submit_practice_and_adapt(
        self,
        learner_id: str,
        outcome: str,
        evidence_submission: str,
        idempotency_key: str,
        updated_hours_per_week: Optional[float] = None,
        proposed_goal_change: Optional[str] = None,
        user_approved_major_change: bool = False,
        disable_memory_retrieval: bool = False,
        simulate_search_failure: bool = False,
    ) -> Dict[str, Any]:
        # G1: Save workflow checkpoint before processing submission
        chk_id = self.db.save_checkpoint(
            learner_id,
            "submit_start",
            {
                "outcome": outcome,
                "idempotency_key": idempotency_key,
                "updated_hours_per_week": updated_hours_per_week,
            },
        )

        # Step 1: Retrieve current state from SQLite (unless disabled for M01-M03 diagnostic run)
        state = retrieve_learner_state(
            self.db,
            self.logger,
            learner_id=learner_id,
            disable_retrieval=disable_memory_retrieval,
        )

        if disable_memory_retrieval or not state.get("current_profile"):
            self.db.complete_checkpoint(chk_id)
            proposal = self.coordinator.recommend_next_action(
                learner_id=learner_id,
                state=state,
                latest_assessment=None,
                resources=[],
                user_context=evidence_submission,
            )
            return {
                "state_recovered": False,
                "proposal": proposal.model_dump(),
            }

        # Step 2 (G8 FIX): Security check against prompt injection attempts inside submitted text.
        # When detected, BLOCK the submission from proceeding — do NOT process tainted evidence.
        if detect_untrusted_instruction(evidence_submission):
            self.db.complete_checkpoint(chk_id)
            self.logger.log_event(
                learner_id,
                "SECURITY_POLICY",
                "detect_untrusted_instruction",
                "blocked_injection_attempt",
                {"submission_excerpt": evidence_submission[:80]},
            )
            return {
                "state_recovered": True,
                "blocked": True,
                "reason": "Submission rejected by security policy: untrusted instruction / prompt injection detected in evidence text. Content is treated as data, not instructions.",
                "preserved_profile": state["current_profile"],
            }

        # Step 3: Check approval rules for profile constraint changes
        current_profile = state["current_profile"]
        if proposed_goal_change:
            needs_approval, msg = check_approval_required(
                current_profile,
                {"goal": proposed_goal_change},
                user_approved=user_approved_major_change,
            )
            if needs_approval:
                self.db.complete_checkpoint(chk_id)
                self.logger.log_event(
                    learner_id, "APPROVAL_GATE", "check_approval_required", "pending_confirmation", {"reason": msg}
                )
                return {
                    "approval_required": True,
                    "message": msg,
                    "preserved_profile": current_profile,
                }

        # Step 4: Apply approved availability correction while preserving historical version
        if updated_hours_per_week is not None and updated_hours_per_week != current_profile["hours_per_week"]:
            old_hours = current_profile["hours_per_week"]
            current_profile = self.db.create_or_update_profile(
                learner_id=learner_id,
                goal=current_profile["goal"],
                current_level=current_profile["current_level"],
                hours_per_week=updated_hours_per_week,
                deadline=current_profile["deadline"],
            )
            state["current_profile"] = current_profile
            self.logger.log_event(
                learner_id,
                "PROFILE_CORRECTION",
                "create_or_update_profile",
                "version_incremented",
                {"new_version": current_profile["version"], "hours_per_week": updated_hours_per_week},
            )
            # G3: Record preference change in long-term episode memory
            self.db.add_episode_summary(
                learner_id=learner_id,
                episode_type="preference",
                summary=f"Weekly study hours corrected from {old_hours} to {updated_hours_per_week} hrs/week (Profile v{current_profile['version']}).",
                metadata={"old_hours": old_hours, "new_hours": updated_hours_per_week, "version": current_profile["version"]},
            )

        active_task = state["active_task"]
        criteria = active_task["acceptance_criteria"] if active_task else ["Loads CSV successfully"]

        # G1: Update checkpoint before evaluation
        self.db.save_checkpoint(
            learner_id,
            "evaluate_evidence",
            {"task_id": active_task["task_id"] if active_task else None, "outcome": outcome},
        )

        # Step 5: Evaluate practice evidence deterministically
        assessment = evaluate_practice_evidence(
            self.logger,
            learner_id=learner_id,
            task_criteria=criteria,
            submission_text=evidence_submission,
            outcome_reported=outcome,
        )

        # Step 6: Record practice session with idempotency protection
        obstacle_desc = "Limited weekly study time / missed session" if outcome == "missed" else None
        practice_record = record_practice_outcome(
            self.db,
            self.logger,
            learner_id=learner_id,
            task_id=active_task["task_id"],
            outcome=outcome,
            evidence_ref=evidence_submission[:200] if evidence_submission else "No evidence submitted",
            assessment=assessment,
            idempotency_key=idempotency_key,
            obstacle_description=obstacle_desc,
        )

        # G3: Record feedback / outcome episode in long-term memory (if not a duplicate retry)
        if not practice_record.get("duplicate_prevented"):
            self.db.add_episode_summary(
                learner_id=learner_id,
                episode_type="feedback",
                summary=(
                    f"Session {practice_record['session_id']} on task '{active_task['title'][:50]}': "
                    f"outcome={outcome}, status={assessment['overall_status']}, score={assessment['score']}."
                ),
                metadata={
                    "session_id": practice_record["session_id"],
                    "overall_status": assessment["overall_status"],
                    "demonstrated": assessment.get("demonstrated_skills", []),
                    "missing_gaps": assessment.get("missing_gaps", []),
                },
            )

        # Refresh state after recording session & obstacle
        state = retrieve_learner_state(self.db, self.logger, learner_id=learner_id)
        resources = search_learning_resources(
            self.db,
            self.logger,
            learner_id=learner_id,
            skill=current_profile["goal"],
            topic="pandas numeric column mean",
            level=current_profile["current_level"],
            simulate_failure=simulate_search_failure,
        )

        # Step 7: Coordinate next decision
        proposal = self.coordinator.recommend_next_action(
            learner_id=learner_id,
            state=state,
            latest_assessment=assessment,
            resources=resources,
            user_context=evidence_submission,
        )

        # Step 8: Validate and commit revised plan if not waiting on clarification or duplicate retry
        revised_plan = state["active_plan"]
        if proposal.decision != "request clarification" and not practice_record.get("duplicate_prevented"):
            revised_plan = revise_learning_plan(
                self.db,
                self.logger,
                learner_id=learner_id,
                proposal=proposal,
                due_date=current_profile["deadline"],
                expected_prior_version=state["active_plan"]["version"] if state["active_plan"] else None,
            )
            # G3: Record new commitment episode in long-term memory
            self.db.add_episode_summary(
                learner_id=learner_id,
                episode_type="commitment",
                summary=f"Plan revised to v{revised_plan['version']} ({proposal.decision}): '{proposal.recommended_task}' ({proposal.estimated_minutes} min).",
                metadata={"plan_id": revised_plan["plan_id"], "version": revised_plan["version"], "decision": proposal.decision},
            )

        # G1: Mark workflow checkpoint completed (clears all in_progress steps for this learner)
        self.db.complete_checkpoint(chk_id)

        return {
            "state_recovered": True,
            "current_profile": current_profile,
            "assessment": assessment,
            "practice_record": practice_record,
            "proposal": proposal.model_dump(),
            "revised_plan": revised_plan,
        }

    # ── G1: Resume interrupted workflow from checkpoint ──

    def resume_from_checkpoint(self, learner_id: str) -> Dict[str, Any]:
        checkpoint = self.db.get_latest_checkpoint(learner_id)
        state = retrieve_learner_state(self.db, self.logger, learner_id=learner_id)
        if not checkpoint:
            return {
                "resumed": False,
                "reason": "No pending in-progress checkpoint found; state is clean.",
                "recovered_state": state,
            }

        self.logger.log_event(
            learner_id,
            "CHECKPOINT_RECOVERY",
            "resume_from_checkpoint",
            "recovered",
            {"checkpoint_id": checkpoint["checkpoint_id"], "step_name": checkpoint["step_name"]},
        )
        self.db.complete_checkpoint(checkpoint["checkpoint_id"])
        return {
            "resumed": True,
            "checkpoint": checkpoint,
            "recovered_state": state,
        }

    # ── G5 & G6: Read-Only Excel Export for Learner / Teacher Review ──

    def export_learner_report(self, learner_id: str, output_path: str) -> str:
        """
        Export learner data to an Excel workbook backed by the SQLite store.
        The spreadsheet is a read-only review view — the SQLite DB remains authoritative.
        Editing the spreadsheet will NOT alter the agent's state or overwrite teacher feedback.
        """
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            raise RuntimeError("openpyxl is required for Excel export. Install with: pip install openpyxl")

        state = self.db.get_learner_state(learner_id)
        profile_history = self.db.get_profile_history(learner_id)
        episodes = self.db.get_episode_summaries(learner_id)

        wb = Workbook()

        # Styling definitions
        header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
        banner_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        banner_font = Font(name="Calibri", bold=True, color="7F6000", size=10)
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        def _style_sheet(ws, headers: List[str], rows: List[List[Any]]) -> None:
            # Row 1: Authoritative store warning banner (G6)
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 1))
            banner_cell = ws.cell(
                row=1,
                column=1,
                value="⚠ READ-ONLY SNAPSHOT — DO NOT EDIT. Exported from authoritative SQLite database. Direct spreadsheet edits are not imported back into agent memory.",
            )
            banner_cell.fill = banner_fill
            banner_cell.font = banner_font
            banner_cell.alignment = Alignment(horizontal="left", vertical="center")

            # Row 2: Column Headers
            for col_idx, h in enumerate(headers, start=1):
                c = ws.cell(row=2, column=col_idx, value=h)
                c.font = header_font
                c.fill = header_fill
                c.alignment = Alignment(horizontal="center", vertical="center")

            # Row 3+: Data rows
            for r_idx, row_vals in enumerate(rows, start=3):
                for c_idx, val in enumerate(row_vals, start=1):
                    cell = ws.cell(row=r_idx, column=c_idx, value=str(val) if val is not None else "")
                    cell.border = thin_border
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

            # Auto-fit column widths (capped at 50)
            for col in ws.columns:
                max_len = max((len(str(cell.value or "")) for cell in col[1:]), default=12)
                col_letter = col[1].column_letter
                ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 50)

        # Sheet 1: Profile & Version History
        ws_prof = wb.active
        ws_prof.title = "Profile History"
        prof_rows = [
            [
                p["version"],
                "ACTIVE" if p["is_current"] else "HISTORICAL",
                p["goal"],
                p["current_level"],
                p["hours_per_week"],
                p["deadline"],
                p["updated_at"],
            ]
            for p in profile_history
        ]
        _style_sheet(
            ws_prof,
            ["Version", "Status", "Goal", "Level", "Hours/Week", "Deadline", "Updated At"],
            prof_rows,
        )

        # Sheet 2: Practice Sessions & Assessments
        ws_sess = wb.create_sheet("Practice Sessions")
        sess_rows = [
            [
                s["session_id"],
                s["task_id"],
                s["outcome"],
                s["assessment_status"],
                s["assessment_summary"].get("score", ""),
                ", ".join(s["assessment_summary"].get("demonstrated_skills", [])),
                ", ".join(s["assessment_summary"].get("missing_gaps", [])),
                s["evidence_ref"],
                s["created_at"],
            ]
            for s in state.get("practice_sessions", [])
        ]
        _style_sheet(
            ws_sess,
            [
                "Session ID",
                "Task ID",
                "Outcome Reported",
                "Assessment Status",
                "Score",
                "Demonstrated Criteria",
                "Missing Gaps",
                "Evidence Excerpt",
                "Timestamp",
            ],
            sess_rows,
        )

        # Sheet 3: Skills
        ws_skills = wb.create_sheet("Skills")
        skill_rows = [
            [sk["skill_name"], sk["mastery_estimate"], sk["evidence_summary"], sk["updated_at"]]
            for sk in state.get("skills", [])
        ]
        _style_sheet(
            ws_skills,
            ["Skill Name", "Mastery Estimate", "Evidence Summary", "Updated At"],
            skill_rows,
        )

        # Sheet 4: Long-Term Episode Summaries (G3)
        ws_ep = wb.create_sheet("Episode Summaries")
        ep_rows = [
            [ep["episode_id"], ep["episode_type"].upper(), ep["summary"], ep["created_at"]]
            for ep in episodes
        ]
        _style_sheet(
            ws_ep,
            ["Episode ID", "Category", "Episode Summary", "Recorded At"],
            ep_rows,
        )

        # Sheet 5: Event Log (Audit Trail)
        ws_events = wb.create_sheet("Event Log")
        event_log = self.logger.get_event_log(learner_id, limit=200)
        event_rows = [
            [
                e["event_id"],
                e.get("run_id", ""),
                e.get("event_type", ""),
                e.get("tool_name", ""),
                e.get("outcome", ""),
                e.get("memory_type", "operational"),
                e.get("timestamp", ""),
            ]
            for e in event_log
        ]
        _style_sheet(
            ws_events,
            ["Event ID", "Run ID", "Event Type", "Tool Name", "Outcome", "Memory Type", "Timestamp"],
            event_rows,
        )

        os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else ".", exist_ok=True)
        wb.save(output_path)

        self.logger.log_event(
            learner_id,
            "REPORT_EXPORT",
            "export_learner_report",
            "success",
            {"output_path": output_path},
        )
        return output_path